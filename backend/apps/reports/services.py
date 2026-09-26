import csv
import json
from datetime import date
from io import BytesIO, StringIO

from django.core.files.uploadedfile import InMemoryUploadedFile
from django.db import transaction
from django.http import HttpResponse

from apps.accounts.models import Account, Investment, Transfer
from apps.analytics.services import analytics_overview, yearly_review
from apps.bills.models import Bill, Subscription
from apps.budgets.models import Budget, BudgetCategory
from apps.core.models import ImportSession, Person, Tag
from apps.core.services.ledger import create_transaction, month_bounds
from apps.core.services.money import money
from apps.debts.models import Debt, DebtPayment, LendingPayment, LendingRecord
from apps.goals.models import GoalContribution, SavingsGoal
from apps.transactions.models import Category, Transaction


FIELD_ALIASES = {
    "date": ["date", "transaction_date", "posted", "value date", "txn date"],
    "amount": ["amount", "value", "debit", "credit", "inflow", "outflow"],
    "description": ["description", "narration", "details", "memo", "particulars"],
    "merchant": ["merchant", "payee", "name"],
    "type": ["type", "transaction_type", "dr/cr", "credit/debit"],
    "category": ["category"],
    "account": ["account"],
    "notes": ["notes", "comment", "remarks"],
}


def _guess_mapping(columns: list[str]) -> dict:
    mapping = {}
    lowered = {col.lower().strip(): col for col in columns}
    for field, aliases in FIELD_ALIASES.items():
        for alias in aliases:
            if alias in lowered:
                mapping[field] = lowered[alias]
                break
    return mapping


def read_tabular(file_obj, name: str) -> list[dict]:
    name = (name or "").lower()
    if name.endswith(".xlsx") or name.endswith(".xls"):
        from openpyxl import load_workbook

        wb = load_workbook(file_obj, read_only=True, data_only=True)
        sheet = wb.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(h).strip() if h is not None else f"col_{i}" for i, h in enumerate(rows[0])]
        data = []
        for row in rows[1:]:
            data.append({headers[i]: row[i] if i < len(row) else None for i in range(len(headers))})
        return data
    text = file_obj.read()
    if isinstance(text, bytes):
        text = text.decode("utf-8-sig")
    reader = csv.DictReader(StringIO(text))
    return list(reader)


def start_import(user, uploaded) -> ImportSession:
    name = getattr(uploaded, "name", "import.csv")
    session = ImportSession.objects.create(user=user, file=uploaded, original_name=name)
    session.file.open("rb")
    rows = read_tabular(session.file, name)
    session.file.close()
    columns = list(rows[0].keys()) if rows else []
    session.detected_columns = columns
    session.mapping = _guess_mapping(columns)
    session.preview = rows[:25]
    session.status = ImportSession.Status.MAPPED
    session.save()
    return session


def _parse_amount(value):
    if value is None or value == "":
        return None
    text = str(value).replace(",", "").replace("Rs.", "").replace("PKR", "").strip()
    if text.startswith("(") and text.endswith(")"):
        text = "-" + text[1:-1]
    return money(text)


def _parse_date(value):
    if value is None or value == "":
        return None
    if isinstance(value, date):
        return value
    text = str(value)[:10]
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y"):
        try:
            from datetime import datetime

            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def preview_import(session: ImportSession, mapping: dict, account: Account) -> list[dict]:
    session.file.open("rb")
    rows = read_tabular(session.file, session.original_name)
    session.file.close()
    preview = []
    for row in rows[:80]:
        amount = _parse_amount(row.get(mapping.get("amount", ""), None))
        tx_date = _parse_date(row.get(mapping.get("date", ""), None))
        description = str(row.get(mapping.get("description", ""), "") or "")
        merchant = str(row.get(mapping.get("merchant", ""), "") or description)
        type_raw = str(row.get(mapping.get("type", ""), "") or "").lower()
        if amount is None:
            continue
        tx_type = "expense"
        if amount < 0:
            amount = money(-amount)
            tx_type = "expense"
        if type_raw in {"income", "credit", "cr", "inflow"}:
            tx_type = "income"
        duplicate = False
        if tx_date:
            duplicate = Transaction.objects.filter(
                user=session.user,
                account=account,
                amount=amount,
                transaction_date=tx_date,
                merchant=merchant,
            ).exists()
        preview.append(
            {
                "date": tx_date.isoformat() if tx_date else None,
                "amount": str(amount),
                "type": tx_type,
                "merchant": merchant,
                "description": description,
                "duplicate": duplicate,
            }
        )
    session.mapping = mapping
    session.preview = preview
    session.status = ImportSession.Status.PREVIEWED
    session.save()
    return preview


@transaction.atomic
def commit_import(session: ImportSession, account: Account, skip_duplicates=True):
    created = 0
    skipped = 0
    for row in session.preview:
        if row.get("duplicate") and skip_duplicates:
            skipped += 1
            continue
        if not row.get("date") or not row.get("amount"):
            skipped += 1
            continue
        create_transaction(
            user=session.user,
            account=account,
            amount=row["amount"],
            transaction_type=row["type"],
            transaction_date=row["date"],
            merchant=row.get("merchant") or "",
            description=row.get("description") or "",
        )
        created += 1
    session.status = ImportSession.Status.IMPORTED
    session.result = {"created": created, "skipped": skipped}
    session.save()
    return session.result


def export_transactions_csv(user, year=None, month=None) -> HttpResponse:
    qs = Transaction.objects.filter(user=user).select_related("account", "category").order_by("transaction_date")
    if year and month:
        start, end = month_bounds(year, month)
        qs = qs.filter(transaction_date__gte=start, transaction_date__lte=end)
    buffer = StringIO()
    writer = csv.writer(buffer)
    writer.writerow(["date", "type", "amount", "merchant", "description", "category", "account", "notes"])
    for tx in qs:
        writer.writerow(
            [
                tx.transaction_date.isoformat(),
                tx.transaction_type,
                str(tx.amount),
                tx.merchant,
                tx.description,
                tx.category.name if tx.category_id else "",
                tx.account.name,
                tx.notes,
            ]
        )
    response = HttpResponse(buffer.getvalue(), content_type="text/csv")
    response["Content-Disposition"] = "attachment; filename=monea-transactions.csv"
    return response


def export_transactions_xlsx(user, year=None, month=None) -> HttpResponse:
    from openpyxl import Workbook

    qs = Transaction.objects.filter(user=user).select_related("account", "category").order_by("transaction_date")
    if year and month:
        start, end = month_bounds(year, month)
        qs = qs.filter(transaction_date__gte=start, transaction_date__lte=end)
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Transactions"
    sheet.append(["date", "type", "amount", "merchant", "description", "category", "account", "notes"])
    for tx in qs:
        sheet.append(
            [
                tx.transaction_date.isoformat(),
                tx.transaction_type,
                float(tx.amount),
                tx.merchant,
                tx.description,
                tx.category.name if tx.category_id else "",
                tx.account.name,
                tx.notes,
            ]
        )
    buffer = BytesIO()
    wb.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = "attachment; filename=monea-transactions.xlsx"
    return response


def build_pdf_report(user, year: int, month: int | None = None) -> HttpResponse:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    if month:
        data = analytics_overview(user, year, month)
        title = date(year, month, 1).strftime("%B %Y") + " Report"
    else:
        data = yearly_review(user, year)
        title = f"{year} Financial Summary"
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=16 * mm)
    styles = getSampleStyleSheet()
    heading = ParagraphStyle("H", parent=styles["Heading1"], fontName="Times-Bold", fontSize=18, textColor=colors.HexColor("#1A1A1A"), spaceAfter=4)
    sub = ParagraphStyle("S", parent=styles["Normal"], textColor=colors.HexColor("#6B6B6B"), fontSize=10, spaceAfter=12)
    body = ParagraphStyle("B", parent=styles["Normal"], fontSize=10, leading=14, textColor=colors.HexColor("#1A1A1A"))
    story = [
        Paragraph("MONEA", sub),
        Paragraph(title, heading),
        Paragraph("Personal Finance OS — private report", sub),
        Spacer(1, 6),
    ]
    if month:
        rows = [
            ["Income", f"Rs. {data['income']:,.0f}"],
            ["Expenses", f"Rs. {data['expenses']:,.0f}"],
            ["Saved", f"Rs. {data['saved']:,.0f}"],
            ["Savings rate", f"{data['savings_rate']}%"],
        ]
    else:
        rows = [
            ["Total income", f"Rs. {data['total_income']:,.0f}"],
            ["Total expenses", f"Rs. {data['total_expenses']:,.0f}"],
            ["Total saved", f"Rs. {data['total_saved']:,.0f}"],
            ["Savings rate", f"{data['savings_rate']}%"],
        ]
    table = Table(rows, colWidths=[90 * mm, 70 * mm])
    table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), "Times-Roman"),
                ("FONTSIZE", (0, 0), (-1, -1), 11),
                ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor("#6B6B6B")),
                ("ALIGN", (1, 0), (1, -1), "RIGHT"),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 8),
                ("LINEBELOW", (0, 0), (-1, -2), 0.3, colors.HexColor("#E7E2DA")),
            ]
        )
    )
    story.append(table)
    story.append(Spacer(1, 14))
    story.append(Paragraph("Category breakdown", heading))
    cats = [["Category", "Amount"]]
    for item in (data.get("categories") or [])[:12]:
        cats.append([item["name"], f"Rs. {money(item['total']):,.0f}"])
    cat_table = Table(cats, colWidths=[90 * mm, 70 * mm])
    cat_table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, 0), "Times-Bold"),
                ("FONTNAME", (0, 1), (-1, -1), "Times-Roman"),
                ("ALIGN", (1, 0), (1, -1), "RIGHT"),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("LINEBELOW", (0, 0), (-1, -1), 0.3, colors.HexColor("#E7E2DA")),
            ]
        )
    )
    story.append(cat_table)
    story.append(Spacer(1, 16))
    story.append(Paragraph("This report is generated from your MONEA records. It is not financial advice.", body))
    doc.build(story)
    response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
    filename = f"monea-{year}-{month or 'year'}.pdf"
    response["Content-Disposition"] = f"attachment; filename={filename}"
    return response


def backup_payload(user) -> dict:
    def dump(qs, fields):
        rows = []
        for obj in qs:
            row = {}
            for field in fields:
                value = getattr(obj, field)
                if hasattr(value, "isoformat"):
                    value = value.isoformat()
                elif hasattr(value, "__str__") and not isinstance(value, (str, int, float, bool, type(None))):
                    value = str(value)
                row[field] = value
            row["id"] = obj.id
            rows.append(row)
        return rows

    return {
        "format": "monea-backup-v1",
        "exported_on": date.today().isoformat(),
        "accounts": dump(
            Account.objects.filter(user=user),
            ["name", "institution", "account_type", "opening_balance", "current_balance", "currency", "color", "icon", "notes"],
        ),
        "categories": dump(Category.objects.filter(user=user), ["name", "icon", "color", "kind", "sort_order"]),
        "transactions": dump(
            Transaction.objects.filter(user=user),
            ["amount", "transaction_type", "merchant", "description", "transaction_date", "notes", "status", "payment_method"],
        ),
        "goals": dump(SavingsGoal.objects.filter(user=user), ["name", "target_amount", "current_amount", "target_date", "notes"]),
        "bills": dump(Bill.objects.filter(user=user), ["name", "amount", "due_date", "frequency", "status"]),
        "subscriptions": dump(Subscription.objects.filter(user=user), ["name", "price", "billing_cycle", "next_billing_date", "status"]),
        "debts": dump(Debt.objects.filter(user=user), ["name", "original_amount", "remaining_amount", "monthly_payment"]),
    }


def restore_backup(user, payload: dict):
    if payload.get("format") != "monea-backup-v1":
        raise ValueError("This file is not a MONEA backup.")
    with transaction.atomic():
        Transaction.objects.filter(user=user).delete()
        Transfer.objects.filter(user=user).delete()
        Account.objects.filter(user=user).delete()
        Category.objects.filter(user=user).delete()
        SavingsGoal.objects.filter(user=user).delete()
        Bill.objects.filter(user=user).delete()
        Subscription.objects.filter(user=user).delete()
        Debt.objects.filter(user=user).delete()
        from apps.core.services.defaults import ensure_user_defaults

        ensure_user_defaults(user)
        name_to_account = {}
        for row in payload.get("accounts", []):
            account = Account.objects.create(
                user=user,
                name=row["name"],
                institution=row.get("institution") or "",
                account_type=row.get("account_type") or Account.Type.OTHER,
                opening_balance=row.get("opening_balance") or 0,
                current_balance=row.get("opening_balance") or 0,
                currency=row.get("currency") or "PKR",
                color=row.get("color") or "#1B7A4E",
                icon=row.get("icon") or "wallet",
                notes=row.get("notes") or "",
            )
            name_to_account[row["name"]] = account
        default_account = next(iter(name_to_account.values()), None)
        if not default_account:
            return
        for row in payload.get("transactions", []):
            create_transaction(
                user=user,
                account=default_account,
                amount=row["amount"],
                transaction_type=row["transaction_type"],
                transaction_date=row["transaction_date"],
                merchant=row.get("merchant") or "",
                description=row.get("description") or "",
                notes=row.get("notes") or "",
            )
